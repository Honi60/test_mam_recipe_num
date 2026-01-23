import os
import json
from datetime import datetime
from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel, 
                             QMessageBox, QFileDialog, QListWidget, QLineEdit, QSpinBox)
from PyQt5.QtGui import QFont
from PyQt5.QtCore import Qt

DB_DIR = r"G:\My Drive\Rentals\RentalsDB"
HISTORY_DIR = os.path.join(DB_DIR, "History")


def load_receipts_by_month_year(month: int, year: int):
    """Load receipts matching the specified month/year from History folder."""
    results = []
    if not os.path.exists(HISTORY_DIR):
        return results
    
    try:
        for filename in sorted(os.listdir(HISTORY_DIR)):
            if filename.endswith('.json'):
                try:
                    fpath = os.path.join(HISTORY_DIR, filename)
                    with open(fpath, 'r', encoding='utf-8') as f:
                        data = json.load(f)
                    # Try to match date field
                    date_str = data.get('Date', '').strip()
                    if date_str:
                        try:
                            date_obj = datetime.strptime(date_str, "%d/%m/%Y")
                            if date_obj.month == month and date_obj.year == year:
                                results.append(data)
                        except ValueError:
                            pass
                except Exception:
                    pass
    except Exception:
        pass
    
    return results


class ToExcelApp_Qt(QWidget):
    def __init__(self):
        super().__init__()
        
        # Setup Hebrew-capable font
        available_fonts = ["Alef", "Segoe UI", "Tahoma"]
        self.hebrew_font = QFont()
        for font_name in available_fonts:
            self.hebrew_font.setFamily(font_name)
            if self.hebrew_font.exactMatch():
                break
        self.hebrew_font.setPointSize(10)
        
        self.init_ui()
    
    def init_ui(self):
        main_layout = QVBoxLayout(self)
        
        # Month and Year selector
        input_layout = QHBoxLayout()
        input_layout.addWidget(QLabel("Month (1-12):"))
        self.month_spinbox = QSpinBox()
        self.month_spinbox.setMinimum(1)
        self.month_spinbox.setMaximum(12)
        self.month_spinbox.setValue(datetime.now().month)
        input_layout.addWidget(self.month_spinbox)
        
        input_layout.addWidget(QLabel("Year:"))
        self.year_spinbox = QSpinBox()
        self.year_spinbox.setMinimum(2020)
        self.year_spinbox.setMaximum(2100)
        self.year_spinbox.setValue(datetime.now().year)
        input_layout.addWidget(self.year_spinbox)
        
        main_layout.addLayout(input_layout)
        
        # List of matching receipts
        main_layout.addWidget(QLabel("Receipts matching month/year:"))
        self.receipt_list = QListWidget()
        main_layout.addWidget(self.receipt_list)
        
        # Button to load receipts
        load_btn = QPushButton("Load Receipts")
        load_btn.clicked.connect(self.load_receipts)
        main_layout.addWidget(load_btn)
        
        # Export button
        export_btn = QPushButton("Export to Excel")
        export_btn.clicked.connect(self.export_to_excel)
        main_layout.addWidget(export_btn)
    
    def load_receipts(self):
        month = self.month_spinbox.value()
        year = self.year_spinbox.value()
        
        receipts = load_receipts_by_month_year(month, year)
        self.receipt_list.clear()
        self.receipts = receipts
        
        if not receipts:
            QMessageBox.information(self, "No Data", f"No receipts found for {month}/{year}")
            return
        
        for data in receipts:
            customer = data.get('customer', 'Unknown')
            date = data.get('Date', 'No date')
            self.receipt_list.addItem(f"{customer} - {date}")
        
        QMessageBox.information(self, "Loaded", f"Found {len(receipts)} receipts for {month}/{year}")
    
    def export_to_excel(self):
        if not hasattr(self, 'receipts') or not self.receipts:
            QMessageBox.warning(self, "Error", "Load receipts first")
            return
        
        output_file = QFileDialog.getSaveFileName(self, "Save Excel File", "", "Excel Files (*.xlsx)")[0]
        if not output_file:
            return
        
        try:
            import openpyxl
        except ImportError:
            QMessageBox.critical(self, "Error", "openpyxl is required. Install with: pip install openpyxl")
            return
        
        try:
            # Create workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            
            # Write header
            headers = ["Customer", "Description", "Invoice No", "Payment", "MAM", "Date", "Bank Account", "Bank", "Check #"]
            for col, header in enumerate(headers, 1):
                ws.cell(row=1, column=col, value=header)
            
            # Write data
            for row_idx, receipt in enumerate(self.receipts, 2):
                ws.cell(row=row_idx, column=1, value=receipt.get('customer', ''))
                ws.cell(row=row_idx, column=2, value=receipt.get('discription', ''))
                ws.cell(row=row_idx, column=3, value=receipt.get('invoice_no', ''))
                ws.cell(row=row_idx, column=4, value=receipt.get('payment', ''))
                ws.cell(row=row_idx, column=5, value=receipt.get('mamVal', ''))
                ws.cell(row=row_idx, column=6, value=receipt.get('Date', ''))
                ws.cell(row=row_idx, column=7, value=receipt.get('bankAccount', ''))
                ws.cell(row=row_idx, column=8, value=receipt.get('BankNumber', ''))
                ws.cell(row=row_idx, column=9, value=receipt.get('CheckNumber', ''))
            
            # Adjust column widths
            for col in range(1, 10):
                ws.column_dimensions[chr(64 + col)].width = 15
            
            wb.save(output_file)
            QMessageBox.information(self, "Success", f"Exported {len(self.receipts)} receipts to {output_file}")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to export: {e}")
