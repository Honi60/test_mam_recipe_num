#!/usr/bin/env python3
"""
Receipt Matrix Report Generator
Creates an Excel table showing all receipts per customer, year, and month.
"""
import json
import os
from pathlib import Path
from collections import defaultdict
import sys
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config.paths import HISTORY_DIR, DB_DIR, CUSTOMERS_FILE, get_data_mode, print_config

OUTPUT_FILE = Path(DB_DIR).parent / "receipt_matrix_report.xlsx"

def load_customers():
    """Load customers from customers_data.json"""
    try:
        with open(CUSTOMERS_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception as e:
        print(f"[ERROR] Error loading customers file: {e}")
        return {}

def parse_date(date_str):
    """Parse date string and return (year, month)"""
    try:
        parts = date_str.split('/')
        if len(parts) == 3:
            day = int(parts[0])
            month = int(parts[1])
            year = int(parts[2])
            if year < 100:
                year += 2000 if year < 50 else 1900
            return (year, month)
    except Exception:
        pass
    return (None, None)

def load_receipts():
    """Load all receipts from history files"""
    receipts_by_customer_month = defaultdict(list)
    
    if not Path(HISTORY_DIR).exists():
        print(f"[ERROR] History directory not found: {HISTORY_DIR}")
        return receipts_by_customer_month
    
    json_files = list(Path(HISTORY_DIR).glob("*.json"))
    
    for file_path in json_files:
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            info_section = data.get('info', {})
            customer_key = info_section.get('customer_name', '')
            
            data_section = data.get('data', {})
            recipe_num = data_section.get('recipeNum', '')
            date_str = data_section.get('Date', '')
            
            year, month = parse_date(date_str)
            
            if customer_key and year and month and recipe_num:
                receipts_by_customer_month[(customer_key, year, month)].append({
                    'recipeNum': recipe_num,
                    'date': date_str
                })
        except Exception as e:
            print(f"[WARNING] Error processing {file_path.name}: {e}")
    
    return receipts_by_customer_month

def get_customer_display_name(customers_data, customer_key):
    """Get the full customer name from customers_data.json"""
    if customer_key in customers_data:
        return customers_data[customer_key].get('customer', customer_key)
    return customer_key

def generate_excel_report(customers_data, receipts_data):
    """Generate Excel report with receipt matrix"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[ERROR] openpyxl is required. Install with: pip install openpyxl")
        return False
    
    # Find all unique years and months
    year_months = set()
    customer_keys = set()
    
    for (customer_key, year, month), receipts in receipts_data.items():
        year_months.add((year, month))
        customer_keys.add(customer_key)
    
    year_months = sorted(year_months)
    all_customers = sorted(set(customers_data.keys()) | customer_keys)
    
    # Find max receipts per customer/month for column sizing
    max_receipts = max((len(r) for r in receipts_data.values()), default=0)
    
    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Receipt Matrix"
    
    # Styles
    header_fill = PatternFill(start_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    missing_fill = PatternFill(start_color="FFC7CE", fill_type="solid")
    missing_font = Font(color="9C0006")
    
    # Write headers
    headers = ["Year", "Month", "Customer Name", "Receipt Count"]
    headers.extend([f"Receipt {i+1}" for i in range(max_receipts)])
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    # Write data rows
    row = 2
    for year, month in year_months:
        for customer_key in all_customers:
            receipts = receipts_data.get((customer_key, year, month), [])
            count = len(receipts)
            customer_name = get_customer_display_name(customers_data, customer_key)
            
            ws.cell(row=row, column=1, value=year)
            ws.cell(row=row, column=2, value=month)
            ws.cell(row=row, column=3, value=customer_name)
            ws.cell(row=row, column=4, value=count)
            
            # Highlight missing receipts
            count_cell = ws.cell(row=row, column=4)
            if count == 0:
                count_cell.fill = missing_fill
                count_cell.font = missing_font
            
            # Write receipt numbers
            for i, receipt in enumerate(receipts):
                ws.cell(row=row, column=5+i, value=receipt['recipeNum'])
            
            row += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 8   # Year
    ws.column_dimensions['B'].width = 8   # Month
    ws.column_dimensions['C'].width = 30  # Customer Name
    ws.column_dimensions['D'].width = 15  # Count
    
    for i in range(max_receipts):
        col = get_column_letter(5 + i)
        ws.column_dimensions[col].width = 12
    
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = 'A2'
    
    try:
        wb.save(OUTPUT_FILE)
        print(f"[OK] Excel report saved to: {OUTPUT_FILE}")
        return True
    except Exception as e:
        print(f"[ERROR] Failed to save Excel file: {e}")
        return False

def main():
    print("Receipt Matrix Report Generator")
    print("=" * 50)
    
    try:
        print_config()
    except UnicodeEncodeError:
        print(f"Data Mode: {get_data_mode()}")
        print(f"Database Dir: {DB_DIR}")
    print(f"Output File: {OUTPUT_FILE}")
    print()
    
    print("Loading customers...")
    customers_data = load_customers()
    print(f"Loaded {len(customers_data)} customers")
    
    print("Loading receipts...")
    receipts_data = load_receipts()
    
    total_receipts = sum(len(r) for r in receipts_data.values())
    print(f"Found {total_receipts} receipts")
    print()
    
    if customers_data or receipts_data:
        print("Generating Excel report...")
        if generate_excel_report(customers_data, receipts_data):
            print()
            print("[OK] Report generated successfully!")
            print(f"Output: {OUTPUT_FILE}")
        else:
            print("[ERROR] Failed to generate report")
    else:
        print("[WARNING] No data found")

if __name__ == "__main__":
    main()
